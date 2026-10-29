from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document

from app.domain.extraction.csv_extractor import extract_text_from_csv
from app.domain.extraction.file_discovery import discover_supported_files
from app.domain.extraction.txt_extractor import extract_text_from_txt
from app.services.document_processing_service import DocumentProcessingService


def test_discover_supported_files_restricts_to_configured_directory(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()

    valid_file = data_dir / "transactions.csv"
    valid_file.write_text("Date,Description,Amount\n2024-02-15,Netflix,12.99\n", encoding="utf-8")

    outside_dir = tmp_path / "outside"
    outside_dir.mkdir()
    (outside_dir / "not_allowed.txt").write_text("This should not be included.", encoding="utf-8")

    discovered = discover_supported_files(data_dir)

    assert [path.name for path in discovered] == ["transactions.csv"]


def test_discover_supported_files_does_not_escape_the_root(tmp_path: Path) -> None:
    data_dir = tmp_path / "allowed"
    data_dir.mkdir()

    nested_dir = data_dir / "nested"
    nested_dir.mkdir()
    (nested_dir / "safe.txt").write_text("inside root", encoding="utf-8")

    outside_dir = tmp_path / "outside"
    outside_dir.mkdir()
    (outside_dir / "escape.csv").write_text("bad", encoding="utf-8")

    discovered = discover_supported_files(data_dir)

    assert [path.name for path in discovered] == ["safe.txt"]
    assert all("outside" not in str(path) for path in discovered)


def test_txt_extractor_reads_synthetic_text(tmp_path: Path) -> None:
    text_file = tmp_path / "statement.txt"
    text_file.write_text("Customer: Jane Doe\nAccount: 1234567890\nAmount: $150.00\n", encoding="utf-8")

    result = extract_text_from_txt(text_file)

    assert "Customer: Jane Doe" in result
    assert "Account: 1234567890" in result


def test_csv_extractor_reads_synthetic_csv(tmp_path: Path) -> None:
    csv_file = tmp_path / "transactions.csv"
    csv_file.write_text(
        "Date,Description,Amount,Balance\n2024-02-15,Netflix,12.99,500.00\n",
        encoding="utf-8",
    )

    result = extract_text_from_csv(csv_file)

    assert "2024-02-15" in result
    assert "Netflix" in result
    assert "12.99" in result


def test_process_directory_marks_extractor_exception_as_failed_without_exposing_sensitive_content(
    tmp_path: Path, monkeypatch
) -> None:
    import app.services.document_processing_service as document_service_module

    data_dir = tmp_path / "data"
    data_dir.mkdir()

    (data_dir / "good.txt").write_text("Customer: Jane Doe\nAmount: $10.00\n", encoding="utf-8")
    (data_dir / "bad.txt").write_text(
        "Phone: +1 415-555-0182\nAccount: 1234567890\n",
        encoding="utf-8",
    )

    original_extractor = document_service_module.extract_text_from_txt

    def raise_on_bad_file(path: str | Path) -> str:
        if Path(path).name == "bad.txt":
            raise RuntimeError("simulated extractor failure")
        return original_extractor(path)

    monkeypatch.setattr(document_service_module, "extract_text_from_txt", raise_on_bad_file)

    result = DocumentProcessingService(data_dir).process_directory()

    assert result["success_count"] == 1
    assert result["failed_count"] == 1
    assert result["safe_to_send"] is False
    assert result["status"] == "unsafe_batch"
    assert "Jane Doe" in result["combined_output"]
    assert "415-555-0182" not in result["combined_output"]
    assert "1234567890" not in str(result["failed_files"])
    assert "simulated extractor failure" not in str(result["failed_files"])
    assert result["failed_files"][0]["file"] == "bad.txt"


def test_pdf_and_docx_extraction_work_with_synthetic_files(tmp_path: Path) -> None:
    from app.domain.extraction.docx_extractor import extract_text_from_docx
    from app.domain.extraction.pdf_extractor import extract_text_from_pdf

    pdf_file = tmp_path / "note.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Customer: Jane Doe\nPhone: +1 415-555-0182\nAmount: $12.34\n")
    doc.save(pdf_file)
    doc.close()

    docx_file = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("Customer: Jane Doe")
    document.add_paragraph("Phone: +1 415-555-0182")
    document.add_paragraph("Amount: $12.34")
    document.save(docx_file)

    assert "Jane Doe" in extract_text_from_pdf(pdf_file)
    assert "415-555-0182" in extract_text_from_docx(docx_file)


def test_process_directory_generates_combined_scrubbed_output(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()

    (data_dir / "statement.txt").write_text(
        "Customer: Jane Doe\nPhone: +1 415-555-0182\nEmail: jane@example.com\nAccount: 1234567890\nDate: 2024-03-17\nDescription: Netflix.com subscription\nAmount: $15.99\nBalance: $1,250.00\n",
        encoding="utf-8",
    )

    (data_dir / "sales.csv").write_text(
        "Date,Description,Amount,Balance\n2024-03-18,Amazon purchase,44.00,1206.00\n",
        encoding="utf-8",
    )

    service = DocumentProcessingService(data_dir)
    result = service.process_directory()

    assert result["success_count"] == 2
    assert result["safe_to_send"] is True
    assert "=== statement.txt ===" in result["combined_output"]
    assert "=== sales.csv ===" in result["combined_output"]
    assert "Jane Doe" in result["combined_output"]
    assert "Netflix.com" in result["combined_output"]
    assert "2024-03-17" in result["combined_output"]
    assert "$15.99" in result["combined_output"]
    assert "jane@example.com" not in result["combined_output"]
    assert "415-555-0182" not in result["combined_output"]
    assert "1234567890" not in result["combined_output"]


def test_process_directory_marks_batch_unsafe_when_any_file_fails(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()

    (data_dir / "good.txt").write_text("Customer: Jane Doe\nAmount: $10.00\n", encoding="utf-8")
    (data_dir / "bad.txt").write_text("", encoding="utf-8")

    service = DocumentProcessingService(data_dir)
    result = service.process_directory()

    assert result["safe_to_send"] is False
    assert result["status"] == "unsafe_batch"
    assert "Jane Doe" in result["combined_output"]
    assert len(result["failed_files"]) == 1


def test_process_directory_rejects_empty_batch(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()

    result = DocumentProcessingService(data_dir).process_directory()

    assert result["safe_to_send"] is False
    assert result["status"] == "empty_batch"
    assert result["combined_output"] == ""


def test_process_directory_does_not_expose_failed_raw_content(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()

    (data_dir / "good.txt").write_text("Customer: Jane Doe\nAmount: $25.00\n", encoding="utf-8")
    (data_dir / "bad.txt").write_text("", encoding="utf-8")

    result = DocumentProcessingService(data_dir).process_directory()

    assert result["safe_to_send"] is False
    assert result["failed_files"][0]["reason"]
    assert "Jane Doe" in result["combined_output"]
    assert "415-555-0182" not in str(result["failed_files"])
    assert "jane@example.com" not in str(result["failed_files"])
